import collections 
import collections.abc
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_presentation():
    prs = Presentation()
    
    # 1. Title Slide
    slide_layout = prs.slide_layouts[0] # Title Slide
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "JAYMANI KRISHI SEWA\nDigital Agritech Marketplace for Farmers & Local Vendors"
    subtitle.text = ("Team Members:\n"
                     "Manish Pandey (2025266283)\n"
                     "Jayant Raj (2025298708)\n\n"
                     "B.Tech CSE Project Proposal / PBL\n"
                     "Department of Computer Science & Engineering\n"
                     "Sharda University")

    # 2. Project Overview
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "2. Project Overview"
    tf = slide.placeholders[1].text_frame
    tf.text = "What the project is:"
    p = tf.add_paragraph()
    p.text = "A digital agritech marketplace bridging the gap between farmers and local vendors."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Problem addressed: Operational friction in agri-retail."
    
    p = tf.add_paragraph()
    p.text = "Target users: Farmers, Local Agricultural Vendors, Administrators."
    
    p = tf.add_paragraph()
    p.text = "Main Purpose:"
    p = tf.add_paragraph()
    p.text = "To provide a modernized ecosystem that empowers local agriculture through digital connectivity, real-time inventory, and digital credit records."
    p.level = 1

    # 3. Problem Statement
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "3. Problem Statement"
    tf = slide.placeholders[1].text_frame
    tf.text = "Existing Problem:"
    
    p = tf.add_paragraph()
    p.text = "Farmers face severe inefficiencies in day-to-day agricultural retail operations."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Limitations of Manual System:"
    p = tf.add_paragraph()
    p.text = "Stock Uncertainty: No real-time visibility into local retail inventory."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Manual Credit Records: Prone to disputes and poor financial tracking."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Fragmented Access: Hard to connect efficiently with available vendors."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Expected Improvement: A unified digital platform to streamline discovery, verify stock, and maintain dispute-free digital ledgers."

    # 4. Objectives and Methodology
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "4. Objectives & Methodology"
    tf = slide.placeholders[1].text_frame
    tf.text = "Project Objectives:"
    p = tf.add_paragraph()
    p.text = "Functional: Enable real-time inventory discovery and digital credit management."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Technical: Build a scalable web architecture using Flask & React."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "User-oriented: Deliver an intuitive interface for non-technical users."
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Methodology (Development Approach):"
    p = tf.add_paragraph()
    p.text = "1. Requirement Analysis & UI/UX Planning\n2. System Design & Database Schema\n3. Module Development (Frontend & Backend)\n4. Integration & API connecting\n5. Testing & Verification\n6. Deployment / Demonstration"
    p.level = 1

    # 5. Tools and Formulas
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "5. Identification of Necessary Tools / Formulas"
    tf = slide.placeholders[1].text_frame
    tf.text = "Technologies Used:"
    p = tf.add_paragraph()
    p.text = "Frontend: React, Tailwind CSS, React Router, HTML5, JSX"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Backend: Python, Flask, Flask-CORS"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Database: SQLite (sqlite3)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Tools: VS Code, Git, Node.js, Python Environment"
    p.level = 1

    p = tf.add_paragraph()
    p.text = "Mathematical Formulas:"
    p = tf.add_paragraph()
    p.text = "No specialized mathematical formula is required; the project primarily uses software/system logic and standard CRUD operations for inventory and credit management."
    p.level = 1

    # 6. Identification of Modules
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "6. Identification of Modules"
    tf = slide.placeholders[1].text_frame
    tf.text = "1. User Authentication & Role Management"
    p = tf.add_paragraph()
    p.text = "Roles: Admin, Farmer, Vendor (Login, Secure access)."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "2. Hyper-Local Discovery (Farmer Dashboard)"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Allows farmers to search and view local products & availability."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "3. Inventory Management (Vendor Dashboard)"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Vendors can add products, manage stock, and view catalog."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "4. Order & Digital Credit System"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Placing orders, tracking status, and maintaining a digital ledger."
    p.level = 1

    # 7. High-Level Design
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "7. High-Level Design"
    tf = slide.placeholders[1].text_frame
    tf.text = "System Architecture:"
    p = tf.add_paragraph()
    p.text = "Client-Server Architecture separating UI from business logic."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Data Flow:"
    p = tf.add_paragraph()
    p.text = "[ Web Client (React + Tailwind) ]"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "        ⬇ (HTTP / JSON via API)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "[ Backend API (Python Flask) ]"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "        ⬇ (SQL Queries)"
    p.level = 1
    p = tf.add_paragraph()
    p.text = "[ Database (SQLite) ]"
    p.level = 1

    # 8. Low-Level Design
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "8. Low-Level Design – Module Wise"
    tf = slide.placeholders[1].text_frame
    tf.text = "Auth Module:"
    p = tf.add_paragraph()
    p.text = "Login component (React) -> POST /api/auth/login -> queries 'users' table."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Discovery Module:"
    p = tf.add_paragraph()
    p.text = "FarmerDashboard -> GET /api/products -> JOIN products & users tables -> Updates UI State."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Order/Credit Module:"
    p = tf.add_paragraph()
    p.text = "orderProduct() -> POST /api/orders -> Inserts into 'orders' table -> Returns Success JSON."
    p.level = 1

    # 9. Working / System Flow
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "9. Working / System Flow"
    tf = slide.placeholders[1].text_frame
    tf.text = "End-to-End Workflow:"
    
    flow = ["User Input (Login via React)",
            "⬇",
            "Dashboard Rendered (React Router based on Role)",
            "⬇",
            "Action Triggered (e.g., View Products / Add Stock)",
            "⬇",
            "API Call to Flask Backend (JSON over HTTP)",
            "⬇",
            "SQLite Database Query/Update",
            "⬇",
            "JSON Response to Client",
            "⬇",
            "Frontend Output (UI Update)"]
    
    for f in flow:
        p = tf.add_paragraph()
        p.text = f
        if "⬇" in f:
            p.level = 1
        else:
            p.level = 0

    # 10. Implementation / Project Evidence
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "10. Implementation / Project Evidence"
    tf = slide.placeholders[1].text_frame
    tf.text = "System Implementation Details:"
    p = tf.add_paragraph()
    p.text = "Frontend App: Built with React, utilizing Vite, React Router, and Tailwind CSS for rapid styling."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Backend Service: RESTful API powered by Flask handling routes like /api/auth/login, /api/products, and /api/orders."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Database Setup: SQLite populated via seed.py with demo data for Admins, Farmers, and Vendors."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Screens/Pages Implemented:"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "Login Page, Farmer Dashboard (Hyper-Local Discovery), Vendor Dashboard (Inventory Management), Admin Dashboard."
    p.level = 1

    # 11. Testing / Verification
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "11. Testing / Verification"
    tf = slide.placeholders[1].text_frame
    tf.text = "Testing Strategies:"
    p = tf.add_paragraph()
    p.text = "Functional Testing: Verified role-based access control (Admin vs Farmer vs Vendor)."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Integration Testing: Verified API connectivity between React frontend and Flask backend using Flask-CORS."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Error Handling: Verified standard HTTP error codes (e.g., 401 for Invalid Credentials)."
    p.level = 1
    
    p = tf.add_paragraph()
    p.text = "Key Test Cases:"
    p.level = 0
    p = tf.add_paragraph()
    p.text = "1. Valid Login -> Redirects to respective Dashboard."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "2. Invalid Login -> Shows 'Invalid credentials' error."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "3. Place Order -> Success message and inserted into DB."
    p.level = 1

    # 12. Expected Outcome
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "12. Expected Outcome"
    tf = slide.placeholders[1].text_frame
    tf.text = "Expected Functionality:"
    p = tf.add_paragraph()
    p.text = "Seamless, unified connection between farmers and their local vendors."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Expected Benefits:"
    p = tf.add_paragraph()
    p.text = "Reduced time searching for supplies and tracking stock."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Dispute-free digital credit tracking instead of manual Bahi-Khata."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Expected User Experience:"
    p = tf.add_paragraph()
    p.text = "Fast, responsive, and intuitive UI suitable for rural/semi-urban populations."
    p.level = 1

    # 13. Future Scope
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "13. Future Scope (Proposed)"
    tf = slide.placeholders[1].text_frame
    tf.text = "Realistic Future Improvements:"
    p = tf.add_paragraph()
    p.text = "Scalability: Migrate backend database from SQLite to MySQL or PostgreSQL for heavy loads."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Language Support: Add Hindi and regional dialect support for broader accessibility."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Digital Payments: Integrate UPI payment gateways for seamless transactions."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Automation: SMS/WhatsApp notifications for order status and credit updates."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "AI Integration: Crop disease detection or price prediction features."
    p.level = 1

    # 14. Communication / Presentation
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "14. Communication & Impact"
    tf = slide.placeholders[1].text_frame
    tf.text = "Key Takeaways:"
    p = tf.add_paragraph()
    p.text = "Targeted Solution: Solves real-world local problems in agricultural retail."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Technical Contribution: Full-stack implementation delivering a modern UI and a robust API."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Social Impact: Empowers the local agricultural economy by organizing fragmented supply chains."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Workflow: Simple and transparent processes for both farmers and vendors."
    p.level = 1

    # 15. Conclusion
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "15. Conclusion"
    tf = slide.placeholders[1].text_frame
    tf.text = "Summary:"
    p = tf.add_paragraph()
    p.text = "Problem: Inefficiencies and stock uncertainty in local agri-retail."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Proposed Solution: Jaymani Krishi Sewa platform."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Core Modules: Auth, Hyper-Local Discovery, Inventory Management, Order/Credit System."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Technology: React (Frontend), Flask (Backend), SQLite (Database)."
    p.level = 1
    p = tf.add_paragraph()
    p.text = "Overall Value: Digitizes and streamlines the unorganized agricultural sector."
    p.level = 1

    # 16. Thank You
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Thank You!"
    subtitle.text = "Questions & Discussion"

    prs.save('JAYMANI_KRISHI_SEWA_Project_Proposal.pptx')

if __name__ == '__main__':
    create_presentation()
